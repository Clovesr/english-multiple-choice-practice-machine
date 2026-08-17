from __future__ import annotations

import io
import sqlite3
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

from fastapi.testclient import TestClient
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject


def _pdf_bytes(text: str = "") -> bytes:
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    if text:
        font = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
        font_ref = writer._add_object(font)
        page[NameObject("/Resources")] = DictionaryObject(
            {
                NameObject("/Font"): DictionaryObject(
                    {NameObject("/F1"): font_ref}
                )
            }
        )
        safe_text = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        stream = DecodedStreamObject()
        stream.set_data(
            f"BT /F1 12 Tf 72 720 Td ({safe_text}) Tj ET".encode("latin-1")
        )
        page[NameObject("/Contents")] = writer._add_object(stream)
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


def _docx_bytes() -> bytes:
    document = Document()
    document.add_heading("Reading Strategies", level=1)
    document.add_paragraph("Skimming identifies the main idea before close reading.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Skill"
    table.cell(0, 1).text = "Purpose"
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


class ResourceApiTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = TemporaryDirectory(ignore_cleanup_errors=True)
        self.database_path = Path(self.temp.name) / "resource-test.db"
        self.database_patch = patch(
            "backend.app.database.DATABASE_PATH", self.database_path
        )
        self.database_patch.start()
        from backend.app.database import initialize_database
        from backend.app.main import app

        initialize_database()
        self.client = TestClient(app)

    def tearDown(self) -> None:
        self.client.close()
        self.database_patch.stop()
        self.temp.cleanup()

    def _import_text(
        self,
        content: str,
        filename: str = "article.txt",
        encoding: str = "utf-8",
    ) -> dict:
        response = self.client.post(
            "/api/resources/import",
            files={"file": (filename, content.encode(encoding), "text/plain")},
        )
        self.assertEqual(response.status_code, 201, response.text)
        return response.json()["resource"]

    def test_gbk_import_duplicate_and_restart_persistence(self) -> None:
        resource = self._import_text(
            "第一段包含中文。\n\nSecond paragraph.",
            filename="中文资料.txt",
            encoding="gbk",
        )
        self.assertEqual(resource["title"], "中文资料")
        self.assertEqual(resource["segment_count"], 2)
        self.assertNotIn("stored_path", resource)

        duplicate = self.client.post(
            "/api/resources/import",
            files={
                "file": (
                    "再次导入.txt",
                    "第一段包含中文。\n\nSecond paragraph.".encode("gbk"),
                    "text/plain",
                )
            },
        )
        self.assertEqual(duplicate.status_code, 409)
        self.assertEqual(duplicate.json()["code"], "duplicate_resource")
        self.assertEqual(duplicate.json()["details"]["existing_id"], resource["id"])

        from backend.app.database import initialize_database

        initialize_database()
        listed = self.client.get("/api/resources").json()
        self.assertEqual(listed["total"], 1)
        self.assertEqual(listed["items"][0]["segment_count"], 2)

    def test_markdown_json_note_and_binary_failure_fallback(self) -> None:
        markdown = self.client.post(
            "/api/resources/import",
            json={
                "title": "Resilience Notes",
                "content": "# Heading\n\nResilience grows.\n\n- deliberate\n- practice",
                "type": "note",
                "format": "md",
            },
        )
        self.assertEqual(markdown.status_code, 201, markdown.text)
        resource = markdown.json()["resource"]
        segments = self.client.get(f"/api/resources/{resource['id']}/segments").json()
        self.assertEqual(
            [item["kind"] for item in segments["items"]],
            ["heading", "paragraph", "paragraph"],
        )

        broken = self.client.post(
            "/api/resources/import",
            files={"file": ("broken.txt", b"\x00\x01\xffbinary", "text/plain")},
        )
        self.assertEqual(broken.status_code, 201, broken.text)
        failed = broken.json()["resource"]
        self.assertEqual(failed["status"], "needs_review")
        self.assertEqual(failed["segment_count"], 0)
        self.assertIn("二进制", failed["parse_error"])
        stored = list((self.database_path.parent / "resources" / failed["uuid"]).glob("original.*"))
        self.assertEqual(len(stored), 1)
        self.assertEqual(stored[0].read_bytes(), b"\x00\x01\xffbinary")

    def test_pdf_docx_and_scanned_pdf_fallback(self) -> None:
        pdf = self.client.post(
            "/api/resources/import",
            files={
                "file": (
                    "reading.pdf",
                    _pdf_bytes("PDF text layer paragraph"),
                    "application/pdf",
                )
            },
        )
        self.assertEqual(pdf.status_code, 201, pdf.text)
        pdf_resource = pdf.json()["resource"]
        self.assertEqual(pdf_resource["format"], "pdf")
        self.assertEqual(pdf_resource["status"], "inbox")
        pdf_segments = self.client.get(
            f"/api/resources/{pdf_resource['id']}/segments"
        ).json()["items"]
        self.assertTrue(
            any("PDF text layer" in item["content"] for item in pdf_segments)
        )
        self.assertEqual(pdf_segments[0]["metadata"], '{"page":1}')

        docx = self.client.post(
            "/api/resources/import",
            files={
                "file": (
                    "strategies.docx",
                    _docx_bytes(),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )
        self.assertEqual(docx.status_code, 201, docx.text)
        docx_resource = docx.json()["resource"]
        self.assertEqual(docx_resource["format"], "docx")
        segments = self.client.get(
            f"/api/resources/{docx_resource['id']}/segments"
        ).json()["items"]
        self.assertEqual(segments[0]["kind"], "heading")
        self.assertIn("Skimming identifies", segments[1]["content"])
        self.assertTrue(any("Skill | Purpose" in item["content"] for item in segments))

        scanned_bytes = _pdf_bytes()
        scanned = self.client.post(
            "/api/resources/import",
            files={"file": ("scan.pdf", scanned_bytes, "application/pdf")},
        )
        self.assertEqual(scanned.status_code, 201, scanned.text)
        scanned_resource = scanned.json()["resource"]
        self.assertEqual(scanned_resource["status"], "needs_review")
        self.assertEqual(scanned_resource["segment_count"], 0)
        self.assertIn("文本层", scanned_resource["parse_error"])
        original = (
            self.database_path.parent
            / "resources"
            / scanned_resource["uuid"]
            / "original.pdf"
        )
        self.assertTrue(original.exists())
        self.assertEqual(original.read_bytes(), scanned_bytes)

    def test_search_progress_update_delete_and_restore(self) -> None:
        resource = self._import_text(
            "Resilience improves with deliberate practice. 韧性来自练习。"
        )
        segments = self.client.get(f"/api/resources/{resource['id']}/segments").json()["items"]
        segment_id = segments[0]["id"]

        english = self.client.get("/api/search", params={"q": "Resilience"})
        self.assertEqual(english.status_code, 200, english.text)
        self.assertEqual(english.json()["total"], 1)
        self.assertIn("<mark>Resilience</mark>", english.json()["items"][0]["snippet"])
        chinese = self.client.get("/api/search", params={"q": "韧性"}).json()
        self.assertEqual(chinese["total"], 1)
        self.assertIn("<mark>韧性</mark>", chinese["items"][0]["snippet"])
        rebuilt = self.client.post("/api/search/rebuild").json()
        self.assertEqual(rebuilt, {"job": "done", "segments": 1})

        progress = self.client.put(
            f"/api/resources/{resource['id']}/progress",
            json={
                "last_segment_id": segment_id,
                "scroll_ratio": 0.5,
                "reading_ms_delta": 30000,
            },
        )
        self.assertEqual(progress.status_code, 200, progress.text)
        self.assertEqual(progress.json()["progress"]["total_reading_ms"], 30000)
        progress = self.client.put(
            f"/api/resources/{resource['id']}/progress",
            json={"scroll_ratio": 0.75, "reading_ms_delta": 5000},
        )
        self.assertEqual(progress.status_code, 200, progress.text)
        self.assertEqual(progress.json()["progress"]["last_segment_id"], segment_id)
        self.assertEqual(progress.json()["progress"]["total_reading_ms"], 35000)
        edited = self.client.put(
            f"/api/resources/{resource['id']}",
            json={"status": "active", "title": "Updated Article"},
        )
        self.assertEqual(edited.json()["title"], "Updated Article")

        deleted = self.client.delete(f"/api/resources/{resource['id']}")
        self.assertEqual(deleted.status_code, 204, deleted.text)
        self.assertEqual(self.client.get("/api/resources").json()["total"], 0)
        trash = self.client.get("/api/trash").json()
        resource_trash = next(item for item in trash if item["resource_type"] == "resource")
        restored = self.client.post(
            f"/api/trash/{resource_trash['id']}/restore", json={}
        )
        self.assertEqual(restored.status_code, 200, restored.text)
        self.assertEqual(self.client.get("/api/resources").json()["total"], 1)

        stored_file = next(
            (self.database_path.parent / "resources" / resource["uuid"]).glob("original.*")
        )
        self.assertTrue(stored_file.exists())
        self.client.delete(f"/api/resources/{resource['id']}")
        trash = self.client.get("/api/trash").json()
        resource_trash = next(
            item
            for item in trash
            if item["resource_type"] == "resource"
            and item["resource_id"] == resource["id"]
        )
        purged = self.client.delete(f"/api/trash/{resource_trash['id']}")
        self.assertEqual(purged.status_code, 200, purged.text)
        self.assertFalse(stored_file.exists())
        self.assertFalse(stored_file.parent.exists())

    def test_purge_refuses_to_delete_outside_resource_storage(self) -> None:
        resource = self._import_text("Storage boundary check.", "boundary.txt")
        outside_file = self.database_path.parent / "do-not-delete.txt"
        outside_file.write_text("keep", encoding="utf-8")
        with sqlite3.connect(self.database_path) as connection:
            connection.execute(
                "UPDATE resources SET stored_path = ? WHERE id = ?",
                (outside_file.name, resource["id"]),
            )

        deleted = self.client.delete(f"/api/resources/{resource['id']}")
        self.assertEqual(deleted.status_code, 204, deleted.text)
        resource_trash = next(
            item
            for item in self.client.get("/api/trash").json()
            if item["resource_type"] == "resource"
        )
        purged = self.client.delete(f"/api/trash/{resource_trash['id']}")
        self.assertEqual(purged.status_code, 200, purged.text)
        self.assertEqual(outside_file.read_text(encoding="utf-8"), "keep")


if __name__ == "__main__":
    unittest.main()
